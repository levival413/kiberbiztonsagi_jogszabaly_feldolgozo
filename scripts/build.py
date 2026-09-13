"""Export the supplied workbook without modifying it. Requires openpyxl."""
from pathlib import Path
import json, re, datetime, shutil
import openpyxl

root = Path(__file__).resolve().parents[1]
source = root / 'docs/downloads/Kiberbiztinsagi_jogszabalyi_feladatok_V.1.0.xlsx'
w = openpyxl.load_workbook(source)
cached = openpyxl.load_workbook(source, data_only=True)
links = []
def cell(c):
    v, href = c.value, None
    if c.data_type == 'f':
        dynamic = re.fullmatch(r'=HYPERLINK\("#\'([^\']+)\'!A"&\(MATCH\("([^"]+)",([^!]+)!\$A\$(\d+):\$A\$(\d+),0\)\+(\d+)\),"([^"]+)"\)', v)
        if dynamic:
            dest, key, lookup, first, last, offset, label = dynamic.groups()
            matches = [i for i,row in enumerate(w[lookup][f'A{first}:A{last}'],1) if row[0].value == key]
            assert len(matches) == 1, (key,matches)
            v = f'=HYPERLINK("#\'{dest}\'!A{matches[0]+int(offset)}","{label}")'
        m = re.fullmatch(r'=HYPERLINK\("((?:[^"]|"")*)","((?:[^"]|"")*)"\)', v)
        if m:
            href, v = [s.replace('""', '"') for s in m.groups()]
            if href.startswith('#'):
                target = re.fullmatch(r"#'([^']+)'!([A-Z]+)(\d+)", href)
                assert target, href
                sn, col, rn = target.groups()
                assert sn in w.sheetnames and int(rn) <= w[sn].max_row
                href = '#' + str(w.sheetnames.index(sn)) + ':' + rn
            links.append(href)
        else:
            m = re.fullmatch(r'=COUNTA\(([^!]+)!([A-Z]+\d+):([A-Z]+\d+)\)', v)
            if m:
                sn, first, last = m.groups()
                v = sum(c.value is not None for row in w[sn][first:last] for c in row)
            else:
                raise ValueError(('Unsupported formula',c.coordinate,v))
    if isinstance(v, (datetime.datetime, datetime.date)):
        v = v.strftime('%Y.%m.%d.')
    result = {'v': '' if v is None else str(v)}
    if href: result['href'] = href
    return result

data = []
for s in w:
    rows = [{'n':row[0].row,'cells':[cell(c) for c in row]} for row in s if any(c.value is not None for c in row)]
    data.append({'name':s.title,'rows':rows})
(root/'docs/data.js').write_text('const WORKBOOK = ' + json.dumps(data,ensure_ascii=False) + ';\n',encoding='utf-8')
print(json.dumps({'sheets':len(data),'rows':{s['name']:len(s['rows']) for s in data},'links':len(links)},ensure_ascii=False))
